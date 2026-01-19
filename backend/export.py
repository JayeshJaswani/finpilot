import io
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

def create_dcf_sheet(wb, dcf, sheet_name, currency_symbol, unit_label):
    ws = wb.create_sheet(title=sheet_name)

    # --- Styles ---
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")
    bold_font = Font(bold=True)
    currency_fmt = f'"{currency_symbol}"#,##0.00'
    percent_fmt = '0.00%'
    
    # Helper to style a header row
    def style_header_row(row_idx, col_start, col_end):
        for c in range(col_start, col_end + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

    # --- 1. Assumptions Section ---
    ws.cell(row=1, column=1, value=f"Assumptions ({unit_label})").font = Font(bold=True, size=14)
    
    # General Inputs
    ws['A3'] = "Historical Revenue"
    ws['B3'] = dcf.historical_revenue
    ws['B3'].number_format = currency_fmt

    ws['A4'] = "Tax Rate"
    ws['B4'] = dcf.tax_rate
    ws['B4'].number_format = percent_fmt

    ws['A5'] = "WACC"
    ws['B5'] = dcf.wacc
    ws['B5'].number_format = percent_fmt

    ws['A6'] = "Terminal Growth Rate"
    ws['B6'] = dcf.terminal_growth_rate
    ws['B6'].number_format = percent_fmt

    ws['A7'] = "Shares Outstanding"
    ws['B7'] = dcf.shares_outstanding
    ws['B7'].number_format = '#,##0'

    ws['A8'] = "Net Debt"
    ws['B8'] = dcf.net_debt
    ws['B8'].number_format = currency_fmt
    
    ws['A9'] = "D&A % of Rev"
    ws['B9'] = dcf.da_percent
    ws['B9'].number_format = percent_fmt

    ws['A10'] = "CapEx % of Rev"
    ws['B10'] = dcf.capex_percent
    ws['B10'].number_format = percent_fmt
    
    ws['A11'] = "NWC % of Rev"
    ws['B11'] = dcf.nwc_percent
    ws['B11'].number_format = percent_fmt

    # Yearly Assumptions Table
    start_row_assumptions = 13
    ws.cell(row=start_row_assumptions, column=1, value="Yearly Assumptions").font = bold_font
    
    headers = ["Metric", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for i, h in enumerate(headers):
        cell = ws.cell(row=start_row_assumptions+1, column=i+1, value=h)
    style_header_row(start_row_assumptions+1, 1, 6)

    # Growth Rates
    ws.cell(row=start_row_assumptions+2, column=1, value="Revenue Growth Rate")
    for i, g in enumerate(dcf.growth_rates):
        cell = ws.cell(row=start_row_assumptions+2, column=i+2, value=g)
        cell.number_format = percent_fmt

    # EBIT Margins
    ws.cell(row=start_row_assumptions+3, column=1, value="EBIT Margin")
    for i, m in enumerate(dcf.ebit_margins):
        cell = ws.cell(row=start_row_assumptions+3, column=i+2, value=m)
        cell.number_format = percent_fmt

    # --- 2. Projections Section ---
    start_row_proj = 18
    ws.cell(row=start_row_proj, column=1, value=f"Projections ({unit_label})").font = Font(bold=True, size=14)
    
    # Headers
    headers = ["Metric", "Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    for i, h in enumerate(headers):
        cell = ws.cell(row=start_row_proj+1, column=i+1, value=h)
    style_header_row(start_row_proj+1, 1, 6)
    
    # --- FORMULAS ---
    # Rows for metrics
    r_rev = start_row_proj + 2
    r_ebit = start_row_proj + 3
    r_tax = start_row_proj + 4
    r_nopat = start_row_proj + 5
    r_da = start_row_proj + 6
    r_capex = start_row_proj + 7
    r_nwc = start_row_proj + 8 # Total NWC
    r_chg_nwc = start_row_proj + 9
    r_ufcf = start_row_proj + 10
    r_df = start_row_proj + 11
    r_pv_ufcf = start_row_proj + 12

    ws.cell(row=r_rev, column=1, value="Revenue")
    ws.cell(row=r_ebit, column=1, value="EBIT")
    ws.cell(row=r_tax, column=1, value="Tax Expense")
    ws.cell(row=r_nopat, column=1, value="NOPAT")
    ws.cell(row=r_da, column=1, value="D&A")
    ws.cell(row=r_capex, column=1, value="CapEx")
    ws.cell(row=r_nwc, column=1, value="Net Working Capital")
    ws.cell(row=r_chg_nwc, column=1, value="Change in NWC")
    ws.cell(row=r_ufcf, column=1, value="Unlevered FCF")
    ws.cell(row=r_df, column=1, value="Discount Factor")
    ws.cell(row=r_pv_ufcf, column=1, value="PV of FCF")

    # Store cell references for valuation
    pv_ufcf_cells = []
    
    for i in range(1, 6): # Years 1 to 5
        col_idx = i + 1
        col_letter = get_column_letter(col_idx)
        prev_col_letter = get_column_letter(col_idx - 1)
        
        # Revenue Formula
        growth_cell = f"{col_letter}{start_row_assumptions+2}"
        margin_cell = f"{col_letter}{start_row_assumptions+3}"
        
        if i == 1:
            rev_formula = f"=$B$3*(1+{growth_cell})"
        else:
            rev_formula = f"={prev_col_letter}{r_rev}*(1+{growth_cell})"
        
        ws.cell(row=r_rev, column=col_idx, value=rev_formula).number_format = currency_fmt

        # EBIT Formula
        ebit_formula = f"={col_letter}{r_rev}*{margin_cell}"
        ws.cell(row=r_ebit, column=col_idx, value=ebit_formula).number_format = currency_fmt

        # Tax Formula
        tax_formula = f"={col_letter}{r_ebit}*$B$4"
        ws.cell(row=r_tax, column=col_idx, value=tax_formula).number_format = currency_fmt

        # NOPAT Formula
        nopat_formula = f"={col_letter}{r_ebit}-{col_letter}{r_tax}"
        ws.cell(row=r_nopat, column=col_idx, value=nopat_formula).number_format = currency_fmt

        # D&A Formula
        da_formula = f"={col_letter}{r_rev}*$B$9"
        ws.cell(row=r_da, column=col_idx, value=da_formula).number_format = currency_fmt

        # CapEx Formula
        capex_formula = f"={col_letter}{r_rev}*$B$10"
        ws.cell(row=r_capex, column=col_idx, value=capex_formula).number_format = currency_fmt
        
        # NWC Total
        nwc_formula = f"={col_letter}{r_rev}*$B$11"
        ws.cell(row=r_nwc, column=col_idx, value=nwc_formula).number_format = currency_fmt

        # Change in NWC
        if i == 1:
             prior_nwc_formula = f"($B$3*$B$11)"
             chg_nwc_formula = f"={col_letter}{r_nwc}-{prior_nwc_formula}"
        else:
             chg_nwc_formula = f"={col_letter}{r_nwc}-{prev_col_letter}{r_nwc}"
        ws.cell(row=r_chg_nwc, column=col_idx, value=chg_nwc_formula).number_format = currency_fmt

        # UFCF Formula
        ufcf_formula = f"={col_letter}{r_nopat}+{col_letter}{r_da}-{col_letter}{r_capex}-{col_letter}{r_chg_nwc}"
        ws.cell(row=r_ufcf, column=col_idx, value=ufcf_formula).number_format = currency_fmt

        # Discount Factor
        df_formula = f"=1/((1+$B$5)^{i})"
        ws.cell(row=r_df, column=col_idx, value=df_formula).number_format = '0.00'

        # PV of UFCF
        pv_formula = f"={col_letter}{r_ufcf}*{col_letter}{r_df}"
        cell = ws.cell(row=r_pv_ufcf, column=col_idx, value=pv_formula)
        cell.number_format = currency_fmt
        pv_ufcf_cells.append(cell.coordinate)

    # --- 3. Valuation Section ---
    r_val = start_row_proj + 14
    ws.cell(row=r_val, column=1, value=f"Valuation ({unit_label})").font = Font(bold=True, size=14)
    
    # Sum of PV UFCF
    ws.cell(row=r_val + 1, column=1, value="Sum of PV of UFCFs")
    sum_pv_formula = f"=SUM({pv_ufcf_cells[0]}:{pv_ufcf_cells[-1]})"
    ws.cell(row=r_val + 1, column=2, value=sum_pv_formula).number_format = currency_fmt
    sum_pv_cell = "B" + str(r_val + 1)

    # Terminal Value
    final_ufcf_cell = get_column_letter(6) + str(r_ufcf)
    tv_formula = f"={final_ufcf_cell}*(1+$B$6)/($B$5-$B$6)"
    
    ws.cell(row=r_val + 2, column=1, value="Terminal Value")
    ws.cell(row=r_val + 2, column=2, value=tv_formula).number_format = currency_fmt
    tv_cell = "B" + str(r_val + 2)

    # PV of Terminal Value
    pv_tv_formula = f"={tv_cell}/((1+$B$5)^5)"
    ws.cell(row=r_val + 3, column=1, value="PV of Terminal Value")
    ws.cell(row=r_val + 3, column=2, value=pv_tv_formula).number_format = currency_fmt
    pv_tv_cell = "B" + str(r_val + 3)

    # Enterprise Value
    ws.cell(row=r_val + 4, column=1, value="Enterprise Value").font = bold_font
    ev_formula = f"={sum_pv_cell}+{pv_tv_cell}"
    ws.cell(row=r_val + 4, column=2, value=ev_formula).font = bold_font
    ws.cell(row=r_val + 4, column=2).number_format = currency_fmt
    ev_cell = "B" + str(r_val + 4)

    # Equity Value
    ws.cell(row=r_val + 5, column=1, value="Less: Net Debt")
    ws.cell(row=r_val + 5, column=2, value="=$B$8").number_format = currency_fmt

    ws.cell(row=r_val + 6, column=1, value="Equity Value").font = bold_font
    equity_formula = f"={ev_cell}-$B$8"
    ws.cell(row=r_val + 6, column=2, value=equity_formula).font = bold_font
    ws.cell(row=r_val + 6, column=2).number_format = currency_fmt
    equity_cell = "B" + str(r_val + 6)

    # Share Price
    ws.cell(row=r_val + 7, column=1, value="Implied Share Price").font = Font(bold=True, size=12, color="008000")
    share_price_formula = f"={equity_cell}/$B$7"
    ws.cell(row=r_val + 7, column=2, value=share_price_formula).font = Font(bold=True, size=12, color="008000")
    ws.cell(row=r_val + 7, column=2).number_format = currency_fmt


def export_dcf_to_excel(dcf, unit_label):
    wb = Workbook()
    # Remove default sheet
    default_ws = wb.active
    wb.remove(default_ws)
    
    # Create USD Sheet Only
    create_dcf_sheet(wb, dcf, "DCF (USD)", "$", unit_label)

    # Buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
